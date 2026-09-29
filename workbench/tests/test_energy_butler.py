import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import app as web
from core import db, paths, repo
from core.config import field_map, fields_by_group
from core.extract import run_extraction, _normalize_llm_value
from core.validate import validate, table_issues
from services import artifacts, survey
from openpyxl import load_workbook


class EnergyButlerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.original = getattr(db._local, 'conn', None)
        db._local.conn = None
        self.patches = [patch.object(db, 'db_path', return_value=str(Path(self.tmp.name)/'test.db')),
                        patch.object(paths, 'OUTPUT_DIR', str(Path(self.tmp.name)/'outputs')),
                        patch.object(paths, 'EXPORT_DIR', str(Path(self.tmp.name)/'exports'))]
        for p in self.patches: p.start()
        self.client = web.app.test_client()
        with self.client.session_transaction() as s: s['dept_key'] = 'dept1'
        self.client.post('/projects/create', data={'name':'能源测试中心','product_type':'energy_butler'})
        self.project = repo.list_projects('dept1')[0]
        self.pp = repo.list_products(self.project['id'])[0]
        self.url = '/p/' + self.pp['id']
        self.expected = json.loads(Path('samples/energy_butler/expected.json').read_text(encoding='utf-8'))

    def tearDown(self):
        if getattr(db._local, 'conn', None): db._local.conn.close()
        db._local.conn = self.original
        for p in reversed(self.patches): p.stop()
        self.tmp.cleanup()

    def save(self, key, value):
        response = self.client.post(self.url+'/field', json={'key':key,'value':value})
        self.assertEqual(response.status_code, 200)
        return response.json

    def fill(self):
        for key, f in field_map('energy_butler').items():
            if not f.get('required'): continue
            v = self.expected.get(key)
            if v is None:
                v = 4 if f['type']=='number' else (f['options'][0]['value'] if f['type']=='select' else ([f['options'][0]['value']] if f['type']=='multiselect' else '现场已确认'))
            self.save(key, v)

    def test_extraction_rows_and_multiline(self):
        text = Path('samples/energy_butler/full.md').read_text(encoding='utf-8')
        result, _ = run_extraction(text, 'energy_butler', provider='rule')
        for k, v in self.expected.items(): self.assertEqual(result[k]['value'], v)
        for key in ['energy_season_combination','energy_chiller_adjust','energy_pump_adjust']:
            self.assertIn('\n', result[key]['value'])
            self.assertIn('西塔', result[key]['value'])
        survey.save_transcript(self.pp, self.project, text)
        survey.extract_latest_source(self.pp, self.project, provider='rule')
        self.save('energy_equipment_list', self.expected['energy_equipment_list'][:1])
        survey.extract_latest_source(self.pp, self.project, provider='rule')
        self.assertEqual(len(json.loads(repo.field_values(self.pp['id'])['energy_equipment_list']['value_json'])), 1)

    def test_workflow_excel_zip_and_gates(self):
        self.fill()
        self.assertTrue(validate(self.pp['id'], 'energy_butler')[0])
        self.assertEqual([g['key'] for g, _ in fields_by_group('energy_butler')], ['basic','docs','property','energy_eligibility','energy_equipment','energy_operation'])
        for page in ['review', 'preview']:
            res = self.client.get(self.url+'/'+page)
            self.assertEqual(res.status_code, 200)
            self.assertIn('西塔机房', res.get_data(as_text=True))
        result = artifacts.generate_artifacts(self.pp, ['equipment_list'])[0]
        self.assertEqual(result['status'], 'success')
        file = result['files'][0]
        self.assertRegex(file['name'], r'能源测试中心_能源管家_设备清单_\d{8}.xlsx')
        wb = load_workbook(file['path']); ws = wb.active
        self.assertEqual(ws.max_column,10); self.assertEqual(ws.max_row,4)
        self.assertEqual(ws['H2'].value,2); self.assertEqual(ws['A4'].value,3)
        self.assertEqual(set(map(str,ws.merged_cells.ranges)), {'B2:B4','C2:C3'})
        wb.close()
        zip_path, _, _ = artifacts.package_artifacts(self.pp)
        with zipfile.ZipFile(zip_path) as z: self.assertTrue(any(n.endswith('.xlsx') for n in z.namelist()))
        file_url = self.url+'/file/equipment_list/'+file['name']
        with self.client.get(file_url) as r: self.assertEqual(r.status_code,200)
        for key, value in [('energy_system_types',['其他系统（风冷机组等）']),('energy_fee_over_1m','否')]:
            state = self.save(key,value)
            self.assertFalse(state['validation_ok'])
            self.assertTrue(any('无法匹配该产品' in i['message'] for i in state['issues']))
            for page in ['review','preview']:
                self.assertIn('无法匹配该产品',self.client.get(self.url+'/'+page).get_data(as_text=True))
            for fn in [lambda:artifacts.generate_artifacts(self.pp,['equipment_list']),lambda:artifacts.regenerate_artifact(self.pp,'equipment_list'),lambda:artifacts.package_artifacts(self.pp)]:
                with self.assertRaises(artifacts.ArtifactServiceError): fn()
            self.assertEqual(self.client.get(file_url).status_code,302)
            self.assertEqual(self.client.get(self.url+'/package/download').status_code,302)
            self.assertTrue(self.save(key,self.expected[key])['validation_ok'])
        saved = repo.field_values(self.pp['id'])['energy_equipment_list']
        self.assertTrue(saved['protected'])
        self.assertEqual(json.loads(saved['value_json']),self.expected['energy_equipment_list'])

    def test_row_validation_and_malformed_input(self):
        self.fill()
        for value in [[{}],[dict(self.expected['energy_equipment_list'][0],quantity='NaN')],[dict(self.expected['energy_equipment_list'][0],quantity=0)],[dict(self.expected['energy_equipment_list'][0],quantity=1.5)],[dict(self.expected['energy_equipment_list'][0],equipment='板式换热器')]]:
            self.assertFalse(self.save('energy_equipment_list',value)['validation_ok'])
        for value in ['text',[1],{}]:
            self.assertEqual(self.client.post(self.url+'/field',json={'key':'energy_equipment_list','value':value}).status_code,400)
        f = field_map('energy_butler')['energy_equipment_list']
        self.assertIsNone(_normalize_llm_value(f, 'not a table'))
        self.assertEqual(_normalize_llm_value(f,[{'equipment':'冷水主机'}])[0]['freq_type'],'待确认')
        self.save('energy_equipment_list',[])
        self.assertFalse(validate(self.pp['id'],'energy_butler')[0])

    def test_negation_and_conflicting_fee_are_not_eligible(self):
        for text in ['水冷机组中央空调，没有冰蓄冷，但有水蓄冷。', '项目采用风冷机组。', 'VRV面板不能被BA控制，BA正常运行。']:
            result, _ = run_extraction(text + '可改造部分的空调电费大于100万。现场冷站调研记录。', 'energy_butler', provider='rule')
            selected = result['energy_system_types']['value'] or []
            self.assertNotIn(self.expected['energy_system_types'][0], selected)
            self.assertNotIn('VRV（末端面板可被 BA 控制且 BA 正常运行）', selected)
        result, _ = run_extraction('可改造部分的空调电费大于100万。可改造部分的空调电费不超过100万。水冷机组无蓄冷。', 'energy_butler', provider='rule')
        self.assertIsNone(result['energy_fee_over_1m']['value'])

    def test_unrelated_yes_does_not_answer_fee(self):
        result, _ = run_extraction('项目是商务中心，水冷机组无蓄冷。冷站供冷区域是办公区，空调电费待确认。', 'energy_butler', provider='rule')
        self.assertIsNone(result['energy_fee_over_1m']['value'])
