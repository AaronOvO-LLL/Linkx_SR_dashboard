# -*- coding: utf-8 -*-
"""设备管家测试：覆盖《设备管家Demo需求文档.md》第 14 章验收场景 A-G。

规则引擎部分是不落库的纯函数测试（需求 12：确定性、可单测）；流程部分沿用
test_fire_butler.py 的临时库夹具，走真实路由与服务层。
"""
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import app as web
from core import db, paths, repo, rules
from core.config import artifact_map, field_map
from core.extract import run_extraction
from core.validate import rules_issues, validate
from services import artifacts, survey

PT = 'device_butler'
TABLE = 'device_room_survey'
SAMPLE_TEXT = ('1 个独立生活水泵房，生活给水分高、中、低 3 个区，3 个生活水泵控制柜，'
               '2 个连通的生活水箱，1 个集水井，1 个集水井泵控制柜；另有 2 个变配电房。')


def output():
    return rules.get_output(PT, 'sensor_bom')


def bom(rows):
    """把录入行换算成 (设备房, 物业设备, 设备数量, 传感器, 传感器数量) 五元组列表。"""
    ds = rules.build_dataset(PT, output(), rows)
    return [tuple(r[c['key']] for c in ds['columns']) for r in ds['rows']]


def room(rows, name, device):
    return next((r for r in rows if r['room_name'] == name and r['device'] == device), None)


class RulesEngineTests(unittest.TestCase):
    """规则引擎：确定性纯函数，不碰数据库。"""

    def test_doc_example_matches_expected_bom(self):
        """需求文档 9.2 的输出示例（连通水箱由用户补为 1 之后）。"""
        rows, _, _ = rules.extract_rows(output(), SAMPLE_TEXT)
        room_row = room(rows, '1#独立生活水泵房', '给水生活水池/水箱')
        room_row['qty'] = 1                      # 场景 F：2 个连通 → 独立水箱数 1
        self.assertEqual(bom(rows), [
            ('1#独立生活水泵房', '给水生活管网始端（高/中/低区）', '3 区', '压力变送器(RS485)', '3'),
            ('1#独立生活水泵房', '给水生活水泵控制柜（高/中/低区）', '3', '串口开关量模块(RS485)', '3'),
            ('1#独立生活水泵房', '给水生活水池/水箱', '1', '水液位传感器(RS485)', '1'),
            ('1#独立生活水泵房', '集水井（污水井）', '1', '水液位传感器(RS485)', '1'),
            ('1#独立生活水泵房', '集水井（污水井）水泵控制柜', '1', '串口开关量模块(RS485)', '1'),
            ('1#独立生活水泵房', '设备房温湿度', '1', '温湿度传感器', '1'),
            ('1#独立生活水泵房', '设备房水浸', '2', '水浸传感器(RS485)', '2'),
            ('1#独立生活水泵房', '数据采集', '1 套', '睿联网关 M1000PRO-A(含电箱+电源)', '1 套'),
            ('1#独立生活水泵房', '设备房摄像头', '1', '摄像头', '1'),
            ('1#变配电房', '设备房温湿度', '1', '温湿度传感器', '1'),
            ('1#变配电房', '设备房水浸', '1', '水浸传感器(RS485)', '1'),
            ('1#变配电房', '数据采集', '1 套', '睿联网关 M1000PRO-A(含电箱+电源)', '1 套'),
            ('1#变配电房', '设备房摄像头', '1', '摄像头', '1'),
            ('2#变配电房', '设备房温湿度', '1', '温湿度传感器', '1'),
            ('2#变配电房', '设备房水浸', '1', '水浸传感器(RS485)', '1'),
            ('2#变配电房', '数据采集', '1 套', '睿联网关 M1000PRO-A(含电箱+电源)', '1 套'),
            ('2#变配电房', '设备房摄像头', '1', '摄像头', '1'),
        ])

    def test_same_type_multiple_rooms_split(self):
        """场景 B / 决策 A1：识别出 2 个就拆成 2 间，命名 1#/2#。"""
        rows, _, _ = rules.extract_rows(output(), '现场有 3 个变配电房，另有 1 个电梯机房。')
        names = list(dict.fromkeys(r['room_name'] for r in rows))
        self.assertEqual(names, ['1#变配电房', '2#变配电房', '3#变配电房', '1#电梯机房'])

    def test_same_room_mentioned_twice_is_not_duplicated(self):
        """同一间房被提到两次只算一间，否则固定项会被补两遍。"""
        rows, _, _ = rules.extract_rows(output(), '另有 1 个空调机房。空调机房没说管网数量。')
        self.assertEqual([r['room_name'] for r in rows].count('1#空调机房'), 5)

    def test_ambiguous_quantity_left_blank_not_guessed(self):
        """场景 C / 决策 A2：连通与否不明 → 留空并降为待确认，不填默认值。"""
        rows, _, conf = rules.extract_rows(output(), SAMPLE_TEXT)
        self.assertEqual(conf, 'low')
        self.assertIsNone(room(rows, '1#独立生活水泵房', '给水生活水池/水箱')['qty'])
        self.assertTrue(any(i['level'] == 'error'
                            for i in rules.build_dataset(PT, output(), rows)['issues']))

    def test_negated_connectivity_cue_is_not_ambiguous(self):
        """「不连通」说明就是 1 个独立水体，不该反过来逼人确认。"""
        rows, _, _ = rules.extract_rows(output(), '1 个独立生活水泵房，生活水箱 2 个，独立设置不连通。')
        self.assertEqual(room(rows, '1#独立生活水泵房', '给水生活水池/水箱')['qty'], 2)

    def test_vague_room_count_marks_pending(self):
        """§13：「几个泵房」→ 待确认，不默认数量。"""
        rows, _, conf = rules.extract_rows(output(), '现场有几个变配电房，具体数量要回去查图纸。')
        self.assertEqual(conf, 'low')
        self.assertEqual(len([r for r in rows if r['room_name'] == '1#变配电房']), 4)

    def test_fixed_items_auto_filled_and_adjustable(self):
        """场景 D / 决策 B3：固定项按空间类型带出默认值，用户可增减。"""
        rows, _, _ = rules.extract_rows(output(), '1 个独立生活水泵房，3 个生活水泵控制柜。')
        self.assertEqual(room(rows, '1#独立生活水泵房', '设备房水浸')['qty'], 2)
        room(rows, '1#独立生活水泵房', '设备房水浸')['qty'] = 3
        self.assertIn(('1#独立生活水泵房', '设备房水浸', '3', '水浸传感器(RS485)', '3'), bom(rows))

    def test_fixed_items_differ_by_space_type(self):
        """需求 8.3：混合消防水泵房摄像头 2、空调机房含冷却水池液位。"""
        rows, _, _ = rules.extract_rows(output(), '1 个混合消防水泵房，另有 1 个空调机房。')
        self.assertEqual(room(rows, '1#混合消防水泵房', '设备房摄像头')['qty'], 2)
        self.assertEqual(room(rows, '1#空调机房', '设备房摄像头')['qty'], 2)
        self.assertEqual(room(rows, '1#空调机房', '冷却水池')['qty'], 1)
        self.assertEqual(room(rows, '1#混合消防水泵房', '设备房水浸')['qty'], 2)

    def test_excluded_items_never_appear(self):
        """场景 E / 决策 C2、C3、微决策 1：排除项不进规则库，也就不可能进清单。"""
        rows, _, _ = rules.extract_rows(
            output(), '2 个变配电房，配电房里有三相电表和变压器。另有 1 个空调机房，冷冻供水管网 2 个。')
        text = json.dumps(bom(rows), ensure_ascii=False)
        for banned in ('三相电表', '漏电流', 'NTC', '管道温度'):
            self.assertNotIn(banned, text)
        self.assertIn(('1#空调机房', '冷冻供水管网', '2 个', '压力变送器(RS485)', '2'), bom(rows))

    def test_optional_pipe_left_blank_produces_no_sensor_and_no_error(self):
        """§13 / 微决策 2：空调机房管网数未提及 → 非必填，留空不配，不报错。"""
        rows, _, _ = rules.extract_rows(
            output(), '1 个空调机房，冷冻供水管网 2 个，冷却供水管网与冷却回水管网现场还没确认。')
        ds = rules.build_dataset(PT, output(), rows)
        self.assertEqual([i['level'] for i in ds['issues']], ['warn', 'warn'])
        self.assertNotIn('冷却供水管网', [r['device'] for r in ds['rows']])

    def test_zero_quantity_produces_no_sensor(self):
        """需求 8.6：设备数量为 0 → 不产出传感器，且不算错误。"""
        rows, _, _ = rules.extract_rows(output(), '1 个独立生活水泵房，3 个生活水泵控制柜。')
        room(rows, '1#独立生活水泵房', '给水生活水泵控制柜（高/中/低区）')['qty'] = 0
        ds = rules.build_dataset(PT, output(), rows)
        self.assertEqual(ds['issues'], [])
        self.assertNotIn('给水生活水泵控制柜（高/中/低区）', [r['device'] for r in ds['rows']])

    def test_unknown_space_type_and_device_are_reported(self):
        """§13：空间类型不在 7 类内、设备不在该空间类型规则内 → 报错，不臆造规则。"""
        rows = [{'room_name': '1#发电机房', 'space_type': '发电机房', 'device': '设备房温湿度', 'qty': 1},
                {'room_name': '1#电梯机房', 'space_type': '电梯机房', 'device': '三相电表', 'qty': 1}]
        messages = [i['message'] for i in rules.build_dataset(PT, output(), rows)['issues']]
        self.assertTrue(any('不在规则库的 7 类取值内' in m for m in messages))
        self.assertTrue(any('不在「电梯机房」的规则内' in m for m in messages))

    def test_invalid_quantity_rejected(self):
        rows = [{'room_name': '1#电梯机房', 'space_type': '电梯机房', 'device': '设备房温湿度', 'qty': -1}]
        self.assertTrue(any('不能为负数' in i['message']
                            for i in rules.build_dataset(PT, output(), rows)['issues']))
        self.assertEqual(rules.parse_qty('3 区'), (3, ''))
        self.assertEqual(rules.parse_qty(''), (None, ''))
        self.assertEqual(rules.parse_qty('abc')[0], None)

    def test_synonyms_normalized_to_rule_library_names(self):
        """决策 6 / §6.2：口语写法归一到规则库标准名，否则换算不出传感器。"""
        rows = [{'room_name': '1#泵房', 'space_type': '生活泵房', 'device': '生活水箱', 'qty': 2}]
        normalized = rules.normalize_rows(output(), rows)
        self.assertEqual(normalized[0]['space_type'], '独立生活水泵房')
        self.assertEqual(normalized[0]['device'], '给水生活水池/水箱')

    def test_materialize_is_idempotent_and_keeps_user_quantity(self):
        """固定项补全幂等，且不覆盖用户已经改过的数量。"""
        rows = [{'room_name': '1#电梯机房', 'space_type': '电梯机房', 'device': '设备房温湿度', 'qty': 1}]
        once = rules.materialize(output(), rows)
        self.assertEqual([r['device'] for r in once],
                         ['设备房温湿度', '设备房水浸', '数据采集', '设备房摄像头'])
        room(once, '1#电梯机房', '设备房水浸')['qty'] = 3
        again = rules.materialize(output(), once)
        self.assertEqual(len(again), 4)
        self.assertEqual(room(again, '1#电梯机房', '设备房水浸')['qty'], 3)

    def test_new_room_gets_fixed_items_on_materialize(self):
        """手动新增一间设备房：口语空间类型被归一，缺的固定项补齐。"""
        rows = [{'room_name': '1#电梯机房', 'space_type': '电梯机房', 'device': '设备房温湿度', 'qty': 1},
                {'room_name': '2#变配电房', 'space_type': '配电房', 'device': '设备房水浸', 'qty': 1}]
        out = rules.materialize(output(), rows)
        added = [r for r in out if r['room_name'] == '2#变配电房']
        self.assertEqual([r['device'] for r in added],
                         ['设备房水浸', '设备房温湿度', '数据采集', '设备房摄像头'])
        self.assertEqual(added[0]['space_type'], '配电房（基础）')

    def test_engine_is_deterministic(self):
        """需求 12：相同输入必得相同输出。"""
        first = bom(rules.extract_rows(output(), SAMPLE_TEXT)[0])
        second = bom(rules.extract_rows(output(), SAMPLE_TEXT)[0])
        self.assertEqual(first, second)

    def test_rules_library_passes_self_check(self):
        self.assertEqual(rules.validate_rules(PT), [])
        self.assertFalse(rules.has_rules('safety_butler'))
        self.assertEqual(rules.validate_rules('safety_butler'), [])

    def test_sample_matches_golden_expectation(self):
        """需求 12：规则库变更须回归——用脱敏样例的期望输出做金标准比对。"""
        expected = json.loads(
            (Path(paths.SAMPLES_DIR) / 'device_butler' / 'expected.json').read_text(encoding='utf-8'))
        self.assertEqual(expected['rules_version'], rules.rules_version(PT))

        sample = (Path(paths.SAMPLES_DIR) / 'device_butler' / 'full.md').read_text(encoding='utf-8')
        result, _ = run_extraction(sample, PT, provider='rule')
        self.assertEqual(result[TABLE]['confidence'], expected['extract_confidence'])
        self.assertEqual(result[TABLE]['value'], expected['extracted_rows'])

        confirmed = json.loads(json.dumps(result[TABLE]['value']))
        for row in confirmed:                       # 决策 B2：连通水箱的独立数由人补
            if row['device'] == '给水生活水池/水箱' and row['qty'] is None:
                row['qty'] = 1
        dataset = rules.build_dataset(PT, output(), confirmed)
        self.assertEqual([c['label'] for c in dataset['columns']], expected['result_columns'])
        self.assertEqual([[r[c['key']] for c in dataset['columns']] for r in dataset['rows']],
                         expected['confirmed_bom'])


class DeviceButlerFlowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.original_conn = getattr(db._local, 'conn', None)
        db._local.conn = None
        self.patches = [
            patch.object(db, 'db_path', return_value=str(Path(self.tmp.name) / 'test.db')),
            patch.object(paths, 'OUTPUT_DIR', str(Path(self.tmp.name) / 'outputs')),
            patch.object(paths, 'EXPORT_DIR', str(Path(self.tmp.name) / 'exports')),
        ]
        for p in self.patches:
            p.start()
        self.client = web.app.test_client()
        with self.client.session_transaction() as session:
            session['dept_key'] = 'dept1'
        self.project = repo.create_project('dept1', {'name': '设备管家测试园区',
                                                     'customer_name': '云澜置业集团有限公司'})
        self.pp = repo.add_product(self.project['id'], PT)

    def tearDown(self):
        if getattr(db._local, 'conn', None):
            db._local.conn.close()
        db._local.conn = self.original_conn
        for p in reversed(self.patches):
            p.stop()
        self.tmp.cleanup()

    def sample(self, name):
        return (Path(paths.SAMPLES_DIR) / 'device_butler' / name).read_text(encoding='utf-8')

    def extract(self, text):
        survey.save_transcript(self.pp, self.project, text)
        return survey.extract_latest_source(self.pp, self.project, provider='rule')

    def table(self):
        return json.loads(repo.field_values(self.pp['id'])[TABLE]['value_json'])

    def save_table(self, rows):
        return survey.save_field_value(self.pp, self.project, {'key': TABLE, 'value': rows})

    def test_scenario_a_full_flow_to_xlsx(self):
        """场景 A：提取 → 计算 → 核对 → 生成 Excel → 下载 → 打包。"""
        page = self.client.get('/projects').get_data(as_text=True)
        self.assertIn('value="device_butler"', page)
        self.extract(self.sample('full.md'))

        values = repo.field_values(self.pp['id'])
        self.assertEqual(values['project_name']['value_json'], json.dumps('云澜国际产业园', ensure_ascii=False))
        row = values[TABLE]
        self.assertEqual(row['status'], 'pending_confirm')       # 连通水箱待确认
        self.assertEqual(row['updated_by'], 'ai')

        # 场景 F：2 个连通的生活水箱 → 用户填独立水箱数 1
        rows = self.table()
        room(rows, '1#独立生活水泵房', '给水生活水池/水箱')['qty'] = 1
        self.save_table(rows)
        self.assertEqual(repo.field_values(self.pp['id'])[TABLE]['protected'], 1)
        ok, issues = validate(self.pp['id'], PT)
        self.assertTrue(ok, [i['message'] for i in issues if i['level'] == 'error'])

        for page_name in ('review', 'preview', 'result'):
            response = self.client.get('/p/%s/%s' % (self.pp['id'], page_name))
            self.assertEqual(response.status_code, 200)

        results = artifacts.generate_artifacts(self.pp, list(artifact_map(PT)))
        self.assertEqual([r['status'] for r in results], ['success'], results)
        produced = results[0]['files'][0]
        self.assertTrue(produced['name'].endswith('.xlsx'))
        self.assertIn('设备管家测试园区', produced['name'])
        self.assertIn('设备管家', produced['name'])
        self.assertGreater(produced['size'], 4000)

        from openpyxl import load_workbook
        sheet = load_workbook(produced['path']).active
        self.assertEqual(sheet.title, '传感器配置清单')
        grid = [[c if c is not None else '' for c in r] for r in sheet.iter_rows(values_only=True)]
        flat = [str(c) for r in grid for c in r]
        self.assertIn('设备房名称', flat)
        self.assertIn('传感器数量', flat)
        self.assertIn('规则库版本：1.0.0', flat)
        body = [r for r in grid if r[0] == '1#独立生活水泵房']
        self.assertEqual(len(body), 9)                          # 5 变动项 + 4 固定项
        self.assertIn(['1#独立生活水泵房', '给水生活管网始端（高/中/低区）', '3 区',
                       '压力变送器(RS485)', '3'], [list(r[:5]) for r in grid])
        self.assertIn(['1#独立生活水泵房', '数据采集', '1 套',
                       '睿联网关 M1000PRO-A(含电箱+电源)', '1 套'], [list(r[:5]) for r in grid])
        self.assertNotIn('三相电表', flat)
        self.assertNotIn('管道温度', flat)

        with self.client.get('/p/%s/file/sensor_bom/%s?mode=dl'
                             % (self.pp['id'], produced['name'])) as response:
            self.assertEqual(response.status_code, 200)

        archive, name, manifest = artifacts.package_artifacts(self.pp)
        self.assertIn('设备管家', name)
        with zipfile.ZipFile(archive) as z:
            self.assertIsNone(z.testzip())
            self.assertTrue(any(n.startswith('01_传感器配置清单/') for n in z.namelist()))
        self.assertEqual(manifest['artifacts'][0]['key'], 'sensor_bom')
        self.assertEqual(manifest['product']['rules_version'], rules.rules_version(PT))

    def test_scenario_b_same_type_multiple_rooms(self):
        """场景 B：「2 个变配电房」输出两组独立行。"""
        self.extract(SAMPLE_TEXT)
        rooms = list(dict.fromkeys(r['room_name'] for r in self.table()))
        self.assertEqual(rooms, ['1#独立生活水泵房', '1#变配电房', '2#变配电房'])

    def test_scenario_c_ambiguous_blocks_generation(self):
        """场景 C：数量模糊 → 待确认，补齐前阻止生成。"""
        self.extract(self.sample('missing.md'))
        self.assertEqual(repo.field_values(self.pp['id'])[TABLE]['status'], 'pending_confirm')
        ok, issues = validate(self.pp['id'], PT)
        self.assertFalse(ok)
        self.assertTrue(any('区数未确定' in i['message'] for i in issues))
        with self.assertRaises(artifacts.ArtifactServiceError):
            artifacts.generate_artifacts(self.pp, ['sensor_bom'])

        rows = self.table()
        for device, qty in (('给水生活管网始端（高/中/低区）', 3),
                            ('给水生活水泵控制柜（高/中/低区）', 3),
                            ('给水生活水池/水箱', 1)):
            room(rows, '1#生活水泵房', device)['qty'] = qty
        self.save_table(rows)
        ok, issues = validate(self.pp['id'], PT)
        self.assertTrue(ok, [i['message'] for i in issues if i['level'] == 'error'])
        self.assertEqual(artifacts.generate_artifacts(self.pp, ['sensor_bom'])[0]['status'], 'success')

    def test_scenario_d_fixed_item_adjustment_flows_to_output(self):
        """场景 D：水浸由 2 增至 3，输出随之更新。"""
        self.extract(SAMPLE_TEXT)
        rows = self.table()
        room(rows, '1#独立生活水泵房', '设备房水浸')['qty'] = 3
        self.save_table(rows)
        self.assertIn(('1#独立生活水泵房', '设备房水浸', '3', '水浸传感器(RS485)', '3'),
                      bom(self.table()))

    def test_manual_save_materializes_and_normalizes(self):
        """手动新增一间设备房：口语空间类型被归一，固定项自动补全。"""
        self.save_table([{'room_name': '1#泵房', 'space_type': '生活泵房',
                          'device': '生活水箱', 'qty': '2'}])
        rows = self.table()
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[0]['space_type'], '独立生活水泵房')
        self.assertEqual(rows[0]['device'], '给水生活水池/水箱')
        self.assertEqual([r['device'] for r in rows[1:]],
                         ['设备房温湿度', '设备房水浸', '数据采集', '设备房摄像头'])

    def test_zeroed_fixed_item_stays_excluded(self):
        """把固定项数量改成 0 = 这间不配；再次保存不会被悄悄改回默认值。"""
        self.extract(SAMPLE_TEXT)
        rows = self.table()
        room(rows, '1#独立生活水泵房', '设备房摄像头')['qty'] = 0
        self.save_table(rows)
        self.save_table(self.table())
        self.assertEqual(room(self.table(), '1#独立生活水泵房', '设备房摄像头')['qty'], 0)
        pumps = [r[1] for r in bom(self.table()) if r[0] == '1#独立生活水泵房']
        self.assertNotIn('设备房摄像头', pumps)
        self.assertIn('设备房水浸', pumps)

    def test_reextraction_protects_human_edits(self):
        """场景 G 前置：人工确认过的表格不被重新提取覆盖。"""
        self.extract(SAMPLE_TEXT)
        rows = self.table()
        room(rows, '1#独立生活水泵房', '设备房水浸')['qty'] = 5
        self.save_table(rows)
        self.extract(SAMPLE_TEXT)
        self.assertEqual(room(self.table(), '1#独立生活水泵房', '设备房水浸')['qty'], 5)

    def test_scenario_g_state_survives_reentry(self):
        """场景 G：退出再进入，文字稿、表格、确认状态与已生成 Excel 可恢复。"""
        text = self.sample('full.md')
        self.extract(text)
        rows = self.table()
        room(rows, '1#独立生活水泵房', '给水生活水池/水箱')['qty'] = 1
        self.save_table(rows)
        artifacts.generate_artifacts(self.pp, ['sensor_bom'])

        self.client.get('/logout')
        self.client.post('/login', data={'dept_key': 'dept1', 'password': '000000'})
        review = self.client.get('/p/%s/review' % self.pp['id'])
        self.assertEqual(review.status_code, 200)
        self.assertIn('1#独立生活水泵房', review.get_data(as_text=True))
        self.assertEqual(repo.latest_source(self.pp['id'])['content'], text)
        self.assertEqual(room(self.table(), '1#独立生活水泵房', '给水生活水池/水箱')['qty'], 1)
        result = self.client.get('/p/%s/result' % self.pp['id'])
        self.assertEqual(result.status_code, 200)
        self.assertIn('传感器配置清单', result.get_data(as_text=True))
        self.assertEqual(repo.artifact_runs(self.pp['id'])[0]['status'], 'success')

    def test_other_products_unaffected(self):
        """规则引擎对没有规则库的产品完全静默，公共流程不受影响。"""
        self.extract(SAMPLE_TEXT)
        other = repo.add_product(self.project['id'], 'fire_butler')
        self.assertEqual(repo.field_values(other['id']), {})
        self.assertFalse(rules.has_rules('fire_butler'))
        self.assertEqual(rules_issues('fire_butler', {}), [])
        self.assertNotIn(TABLE, field_map('fire_butler'))

    def test_low_relevance_transcript_rejected(self):
        """§13：与设备管家无关的文字稿被关联度检查拦下。"""
        from core.extract import ExtractionError
        with self.assertRaisesRegex(ExtractionError, '设备管家'):
            run_extraction('今天天气晴朗，我们准备去河边走走再回家吃晚饭，顺便聊聊周末的安排。',
                           PT, provider='rule')


if __name__ == '__main__':
    unittest.main()
