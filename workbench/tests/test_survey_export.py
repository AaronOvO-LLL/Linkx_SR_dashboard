"""全产品完整调研 Excel 的下载、数据完整性与项目权限回归。"""
import json
import unittest
from io import BytesIO
from contextlib import closing
from pathlib import Path
from urllib.parse import unquote
from unittest.mock import patch

from openpyxl import load_workbook

import app as web
from core import paths, repo
from core.config import field_map, list_products
from services import survey_export
from tests import test_input_modes


class SurveyExportTests(unittest.TestCase):
    def setUp(self):
        self.fixture = test_input_modes.InputModeTests()
        self.fixture.setUp()
        self.client = self.fixture.client
        self.project = self.fixture.project

    def tearDown(self):
        self.fixture.tearDown()

    def product(self, product_type):
        return repo.ensure_product(self.project['id'], product_type)

    def save(self, pp, key, value):
        response = self.client.post('/p/%s/field' % pp['id'], json={'key': key, 'value': value})
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))

    def download(self, pp):
        with self.client.get('/p/%s/survey.xlsx' % pp['id']) as response:
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.mimetype, survey_export.MIME_TYPE)
            self.assertIn('attachment;', response.headers['Content-Disposition'])
            self.assertIn('完整调研结果', unquote(response.headers['Content-Disposition']))
            self.assertEqual(response.headers['Cache-Control'], 'no-store')
            return load_workbook(BytesIO(response.data))

    def answers(self, book):
        return {row[1]: row for row in book['调研结果'].iter_rows(min_row=2, values_only=True)}

    def test_all_products_have_entry_and_download_before_generation(self):
        products = list_products()
        self.assertEqual(len(products), 5)
        for product in products:
            with self.subTest(product=product['product_type']):
                pp = self.product(product['product_type'])
                project_html = self.client.get('/projects/%s' % self.project['id']).get_data(as_text=True)
                self.assertIn('/p/%s/result' % pp['id'], project_html)
                html = self.client.get('/p/%s/result' % pp['id']).get_data(as_text=True)
                self.assertIn('下载完整调研 Excel', html)
                self.assertIn('/p/%s/survey.xlsx' % pp['id'], html)
                self.assertIn('还没有生成任何材料', html)
                with closing(self.download(pp)) as book:
                    fields = field_map(pp['product_type'])
                    self.assertEqual(book['调研结果'].max_row, len(fields) + 1)
                    self.assertEqual(set(self.answers(book)), {f['label'] for f in fields.values()})
                    self.assertIn('（未填写）', [r[2] for r in self.answers(book).values()])
                    self.assertIn('校验与提示', book.sheetnames)
                    for field in fields.values():
                        if field['type'] == 'table':
                            self.assertIn(field['label'], book.sheetnames)
                self.assertEqual(repo.artifact_runs(pp['id']), [])
                self.assertIsNone(repo.latest_export(pp['id']))

    def test_scalar_answers_keep_zero_multiselect_multiline_and_formula_text(self):
        pp = self.product('asset_butler')
        text = '=HYPERLINK("https://example.invalid","原样文本")\n第二行'
        for key, value in {'asset_customer_name': text, 'operating_project_count': 0,
                           'asset_types': ['产业园区', '写字楼'],
                           'submitter_contact': '陈岚 13800000000'}.items():
            self.save(pp, key, value)
        with closing(self.download(pp)) as book:
            answers = self.answers(book)
            fields = field_map('asset_butler')
            self.assertEqual(answers[fields['asset_customer_name']['label']][2], text)
            self.assertEqual(answers[fields['operating_project_count']['label']][2], 0)
            self.assertEqual(answers[fields['asset_types']['label']][2], '产业园区、写字楼')
            self.assertEqual(answers[fields['submitter_contact']['label']][4], '已人工确认')
            cell = next(cell for row in book['调研结果'] for cell in row if cell.value == text)
            self.assertEqual(cell.data_type, 's')
            self.assertTrue(cell.alignment.wrap_text)
            self.assertEqual(book['调研结果'].freeze_panes, 'A2')

    def test_energy_export_includes_all_fields_and_unmerged_raw_equipment(self):
        pp = self.product('energy_butler')
        sample = json.loads((Path(paths.SAMPLES_DIR) / 'energy_butler/expected.json').read_text(encoding='utf-8'))
        for key, value in sample.items():
            self.save(pp, key, value)
        operation = '夏季：东塔两台主机运行\n冬季：西塔一台主机运行'
        self.save(pp, 'energy_season_combination', operation)
        with closing(self.download(pp)) as book:
            field = field_map('energy_butler')['energy_equipment_list']
            sheet = book[field['label']]
            self.assertEqual(sheet.max_row, len(sample['energy_equipment_list']) + 1)
            self.assertEqual(list(next(sheet.values)), ['序号'] + [c['label'] for c in field['columns']])
            for index, original in enumerate(sample['energy_equipment_list'], 2):
                for column, definition in enumerate(field['columns'], 2):
                    expected = original.get(definition['key'])
                    if definition.get('type') == 'number' and expected not in (None, ''):
                        expected = float(expected)
                    self.assertEqual(sheet.cell(index, column).value, None if expected == '' else expected)
            self.assertEqual(list(sheet.merged_cells.ranges), [])
            answers = self.answers(book)
            self.assertEqual(answers[field_map('energy_butler')['energy_season_combination']['label']][2],
                             operation)

    def test_device_export_includes_raw_rows_computed_sensors_and_errors(self):
        pp = self.product('device_butler')
        sample = json.loads((Path(paths.SAMPLES_DIR) / 'device_butler/expected.json').read_text(encoding='utf-8'))
        self.save(pp, 'device_room_survey', sample['extracted_rows'])
        with closing(self.download(pp)) as book:
            field = field_map('device_butler')['device_room_survey']
            self.assertEqual(book[field['label']].max_column, len(field['columns']) + 1)
            self.assertGreater(book['传感器配置清单'].max_row, 1)
            messages = [r[2] for r in book['校验与提示'].iter_rows(min_row=2, values_only=True)]
            self.assertTrue(any('未确定' in message for message in messages))
            # 不完整数量不会被导成 0，也不会让同一行的原始设备信息消失。
            raw = list(book[field['label']].values)
            self.assertIsNone(raw[3][-1])
            self.assertIn('给水生活水池/水箱', raw[3])

    def test_each_download_reads_latest_shared_and_product_answers(self):
        fire = self.product('fire_butler')
        safety = self.product('safety_butler')
        self.save(fire, 'project_name', '公共项目第一版')
        self.save(fire, 'fire_renovation_cost', 10)
        with closing(self.download(fire)) as book:
            self.assertEqual(self.answers(book)['项目名称'][2], '公共项目第一版')
        self.save(safety, 'project_name', '公共项目第二版')
        self.save(fire, 'fire_renovation_cost', 20)
        with closing(self.download(fire)) as book:
            answers = self.answers(book)
            self.assertEqual(answers['项目名称'][2], '公共项目第二版')
            self.assertEqual(answers[field_map('fire_butler')['fire_renovation_cost']['label']][2], 20)
        with closing(self.download(safety)) as book:
            self.assertNotIn(field_map('fire_butler')['fire_renovation_cost']['label'], self.answers(book))

    def test_download_authorizes_before_export_and_unknown_id_is_404(self):
        pp = self.product('asset_butler')
        url = '/p/%s/survey.xlsx' % pp['id']
        anonymous = web.app.test_client()
        outsider = web.app.test_client()
        with outsider.session_transaction() as session:
            session['dept_key'] = 'dept2'
        with patch.object(web.survey_export, 'export_survey') as export:
            self.assertEqual(anonymous.get(url).status_code, 302)
            self.assertEqual(outsider.get(url).status_code, 403)
            self.assertEqual(self.client.get('/p/missing/survey.xlsx').status_code, 404)
            export.assert_not_called()

    def test_sheet_names_and_download_name_are_safe(self):
        pp = self.product('asset_butler')
        repo.update_project(self.project['id'], {'name': '客户/项目:*?\r\n测试'})
        stream, filename = survey_export.export_survey(pp)
        self.assertNotRegex(filename, r'[<>:"/\\|?*\x00-\x1f]')
        with closing(load_workbook(stream)) as book:
            self.assertIn('客户/项目:*?\n测试', [r[1] for r in book['导出说明'].values])
        book = survey_export.Workbook()
        book.remove(book.active)
        for _ in range(2):
            survey_export._sheet(book, '明细/[]:*?' * 8, ['列'], [20])
        self.assertEqual(len(set(book.sheetnames)), 2)
        self.assertTrue(all(len(name) <= 31 for name in book.sheetnames))


if __name__ == '__main__':
    unittest.main()

