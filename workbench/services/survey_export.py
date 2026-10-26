"""按产品字段模型导出完整调研快照，不依赖材料是否已生成。"""
import math
import re
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import generate
from core.capabilities import require_capability
from core.validate import is_empty, validate


MIME_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _cell_value(value):
    if isinstance(value, (list, dict)):
        value = generate._value_to_text(value)
    if isinstance(value, float) and not math.isfinite(value):
        value = str(value)
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub('', value.replace('\r\n', '\n').replace('\r', '\n'))
    return value


def _append(sheet, values):
    """所有外部文本显式写为字符串，包括以等号开头的答案。"""
    sheet.append([_cell_value(value) for value in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def _number(value, field):
    if field.get('type') == 'number' and isinstance(value, str):
        try:
            number = float(value)
            if math.isfinite(number):
                return int(number) if number.is_integer() else number
        except ValueError:
            pass
    return value


def _sheet(book, title, headers, widths):
    title = re.sub(r'[\\/*?:\[\]\x00-\x1f]', '_', title).strip(" '") or '明细'
    base, name, suffix = title[:31], title[:31], 1
    while name.lower() in {item.lower() for item in book.sheetnames}:
        suffix += 1
        tail = '_%d' % suffix
        name = base[:31 - len(tail)] + tail
    sheet = book.create_sheet(name)
    _append(sheet, headers)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def _style(book):
    for sheet in book:
        for row in sheet:
            for cell in row:
                cell.font = Font(name='微软雅黑', size=11, color='1F2937')
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='264B65')
            cell.font = Font(name='微软雅黑', bold=True, size=11, color='FFFFFF')
        sheet.row_dimensions[1].height = 30
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions


def export_survey(pp):
    """读取当前已保存数据，返回内存中的 xlsx 和包含项目、产品、时间的文件名。

    未完成的调研也可导出；核对状态、校验问题和计算异常随文件一起提供，
    避免把草稿或不完整的计算清单误当作已确认材料。
    """
    require_capability(pp['product_type'], 'artifact_generation')
    context = generate.build_context(pp['id'], pp['product_type'])
    valid, issues = validate(pp['id'], pp['product_type'])
    stamp = datetime.now(timezone(timedelta(hours=8)))
    book = Workbook()
    book.remove(book.active)
    overview = _sheet(book, '导出说明', ['项目', '内容'], [24, 95])
    for label, value in [
        ('项目记录名称', context['project']['name']),
        ('产品', context['product']['name']),
        ('导出时间（北京时间）', stamp.strftime('%Y-%m-%d %H:%M:%S')),
        ('字段模板版本', context['meta']['field_template_version']),
        ('规则库版本', context['meta']['rules_version'] or '不适用'),
        ('校验结果', '通过' if valid else '存在待补充或待处理问题，详见「校验与提示」'),
        ('数据范围', '当前产品全部已保存的调研字段、明细表格和计算清单。'),
        ('使用说明', '本文件是导出时的调研快照，不代表全部答案已人工确认；'
         '请结合核对状态和校验提示使用。网页中尚未保存的修改不包含在内。'),
    ]:
        _append(overview, [label, value])

    results = _sheet(book, '调研结果', ['分组', '调研项', '结果', '单位', '核对状态'],
                     [22, 34, 85, 12, 20])
    for group in context['groups']:
        for field in group['fields']:
            value = field['value']
            if field['type'] == 'table' and (isinstance(value, list) or is_empty(value)):
                value = value or []
                columns = field['columns']
                detail = _sheet(book, field['label'],
                                ['序号'] + [col['label'] for col in columns],
                                [8] + [26] * len(columns))
                for index, row in enumerate(value, 1):
                    _append(detail, [index] + [_number(row.get(col['key']), col) for col in columns])
                if not value:
                    _append(detail, ['（未填写）'])
                value = '%d 行，详见「%s」' % (len(value), detail.title)
            elif field['type'] == 'attachment' and value:
                value = field['text']
            elif is_empty(value):
                value = '（未填写）'
            else:
                value = _number(value, field)
            _append(results, [group['name'], field['label'], value,
                              field['unit'], field['status_label']])

    for dataset in context['datasets'].values():
        columns = dataset['columns']
        sheet = _sheet(book, dataset['label'] or dataset['key'],
                       [col['label'] for col in columns], [26] * len(columns))
        for row in dataset['rows']:
            _append(sheet, [row.get(col['key']) for col in columns])
        if not dataset['rows']:
            _append(sheet, ['（暂无计算结果，请结合校验与提示检查原始调研数据）'])

    if issues:
        sheet = _sheet(book, '校验与提示', ['级别', '调研项', '问题说明'], [14, 34, 95])
        for issue in issues:
            field = context['detail'].get(issue.get('field'), {})
            _append(sheet, ['错误' if issue['level'] == 'error' else '提示',
                            field.get('label', issue.get('field', '')),
                            issue['message']])

    _style(book)
    stream = BytesIO()
    book.save(stream)
    book.close()
    stream.seek(0)
    # 只作为下载名使用，不落盘；同时清除 Windows 非法字符与 HTTP 控制字符。
    name = '%s_%s_完整调研结果_%s.xlsx' % (
        re.sub(r'[<>:"/\\|?*\x00-\x1f\x7f]', '_', context['project']['name'])[:60].strip('. ') or '项目',
        re.sub(r'[<>:"/\\|?*\x00-\x1f\x7f]', '_', context['product']['name'])[:40],
        stamp.strftime('%Y%m%d_%H%M%S'))
    return stream, name
